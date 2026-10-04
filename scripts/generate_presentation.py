from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
APP = ROOT / "Dehqon-AI" / "artifacts" / "dehqon-ai"
OUTPUT = ROOT / "Taqdimot_Dehqon_AI.pptx"
LOGO = APP / "assets" / "images" / "icon.png"

GREEN = "0D3B2A"
PRIMARY = "2F7D4B"
YELLOW = "F6D36B"
CREAM = "F6F8F1"
INK = "16352A"
MUTED = "66766A"
SAGE = "E5F1E5"
WHITE = "FFFFFF"
LINE = "DCE7DA"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_box(slide, x, y, w, h, fill, rounded=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, text, x, y, w, h, size=16, fill=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.01)
    frame.margin_right = Inches(0.01)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    return shape


def add_footer(slide, number, dark=False):
    tone = "B8CCBC" if dark else MUTED
    add_text(slide, "DEHQON AI  /  LOYIHA TAQDIMOTI", 0.65, 7.12, 5.5, 0.18, 8, tone, True)
    add_text(slide, f"{number:02d}", 12.05, 7.08, 0.6, 0.22, 9, tone, True, align=PP_ALIGN.RIGHT)


def add_heading(slide, eyebrow, title, number, dark=False):
    title_color = WHITE if dark else INK
    eyebrow_color = YELLOW if dark else PRIMARY
    add_text(slide, eyebrow.upper(), 0.72, 0.48, 10.8, 0.25, 9, eyebrow_color, True)
    add_text(slide, title, 0.72, 0.88, 11.8, 0.7, 27, title_color, True, "Aptos Display")
    add_footer(slide, number, dark)


def add_card(slide, x, y, w, h, title, body, number=None, accent=PRIMARY):
    add_box(slide, x, y, w, h, WHITE, rounded=True, line=LINE)
    add_box(slide, x, y, 0.08, h, accent)
    if number:
        add_text(slide, number, x + 0.25, y + 0.22, 0.48, 0.3, 11, accent, True)
        title_y = y + 0.64
    else:
        title_y = y + 0.3
    add_text(slide, title, x + 0.25, title_y, w - 0.5, 0.42, 16, INK, True, "Aptos Display")
    add_text(slide, body, x + 0.25, title_y + 0.52, w - 0.5, h - (title_y - y) - 0.65, 11, MUTED)


# Cover
slide = prs.slides.add_slide(blank)
add_box(slide, 0, 0, 13.333, 7.5, GREEN)
add_box(slide, 0.7, 0.68, 0.1, 5.95, YELLOW)
add_text(slide, "RAQAMLI QISHLOQ XO‘JALIGI", 1.05, 0.92, 5.8, 0.3, 10, YELLOW, True)
add_text(slide, "Dehqon AI", 1.02, 1.72, 7.4, 1.12, 40, WHITE, True, "Aptos Display")
add_text(slide, "Daladagi o‘zbekcha\nraqamli yordamchi", 1.06, 3.0, 7.3, 1.35, 27, "E3EFE3", True, "Aptos Display")
add_text(slide, "O‘simlik tahlili, ob-havo, sug‘orish va subsidiya ma’lumotlari — bitta mobil ilovada.",
         1.08, 4.8, 6.8, 0.75, 15, "D0DFD2")
slide.shapes.add_picture(str(LOGO), Inches(8.85), Inches(1.3), width=Inches(3.25), height=Inches(3.25))
add_text(slide, "LOYIHA TAQDIMOTI  ·  2026", 1.08, 6.82, 5.5, 0.22, 9, "B8CCBC", True)
add_footer(slide, 1, True)

# Problem
slide = prs.slides.add_slide(blank)
add_box(slide, 0, 0, 13.333, 7.5, CREAM)
add_heading(slide, "01 / Muammo", "Fermer qarori ko‘p ma’lumotga tayanadi", 2)
add_text(slide, "Kasallik alomatini baholash, sug‘orishni rejalash va yordam dasturini topish turli manbalarni talab qiladi.",
         0.74, 1.78, 10.9, 0.56, 15, MUTED)
add_card(slide, 0.72, 2.72, 3.82, 2.4, "Kasallikni erta payqash", "Bargdagi belgini to‘g‘ri talqin qilish va keyingi qadamni tanlash har doim ham oson emas.", "A", PRIMARY)
add_card(slide, 4.76, 2.72, 3.82, 2.4, "Suvdan oqilona foydalanish", "Maydon, ekin va sug‘orish usulini hisobga olgan taxminiy reja kerak.", "B", "4E8A76")
add_card(slide, 8.8, 2.72, 3.82, 2.4, "Ma’lumot tarqoqligi", "Ob-havo, tavsiya va subsidiya yo‘riqnomalarini alohida qidirishga to‘g‘ri keladi.", "C", "C9851B")
add_text(slide, "Maqsad: muhim boshlang‘ich ma’lumotlarni sodda, o‘zbekcha interfeysda jamlash.",
         0.75, 5.65, 11.7, 0.5, 16, INK, True)

# Solution
slide = prs.slides.add_slide(blank)
add_box(slide, 0, 0, 13.333, 7.5, WHITE)
add_heading(slide, "02 / Yechim", "Rasm va maydon ma’lumotidan amaliy tavsiyagacha", 3)
steps = [
    ("01", "Kiritish", "Barg rasmini tanlash yoki ekin, tuproq va maydonni belgilash."),
    ("02", "Tahlil", "Mobil qurilmadagi TFLite model rasmni tekshiradi; sug‘orish hisobi mavjud me’yorlardan foydalanadi."),
    ("03", "Keyingi qadam", "Natija, ehtiyotkorlik eslatmasi va umumiy tavsiyalar ilovada ko‘rsatiladi."),
]
for idx, (num, title, body) in enumerate(steps):
    x = 0.75 + idx * 4.15
    add_box(slide, x, 2.05, 3.7, 2.6, CREAM, rounded=True)
    add_text(slide, num, x + 0.25, 2.32, 0.65, 0.4, 18, PRIMARY, True, "Aptos Display")
    add_text(slide, title, x + 0.25, 2.95, 3.1, 0.4, 18, INK, True, "Aptos Display")
    add_text(slide, body, x + 0.25, 3.55, 3.15, 0.8, 12, MUTED)
    if idx < 2:
        add_text(slide, ">", x + 3.75, 3.0, 0.35, 0.4, 20, YELLOW, True, align=PP_ALIGN.CENTER)
add_box(slide, 0.75, 5.25, 11.85, 0.85, GREEN, rounded=True)
add_text(slide, "Mobil tahlil internet bo‘lmaganda ham ishlashi uchun model qurilmaga joylangan.",
         1.05, 5.52, 11.2, 0.3, 14, WHITE, True)

# Features
slide = prs.slides.add_slide(blank)
add_box(slide, 0, 0, 13.333, 7.5, CREAM)
add_heading(slide, "03 / Imkoniyatlar", "Bir ilovada kundalik dehqonchilik yordamchilari", 4)
add_card(slide, 0.72, 1.95, 5.75, 1.75, "O‘simlik tahlili", "Rasm asosida 38 sinfli model taxmin beradi; ishonchsiz natijada qayta suratga olish tavsiya qilinadi.", accent=PRIMARY)
add_card(slide, 6.78, 1.95, 5.75, 1.75, "Sug‘orish", "Ekin, tuproq holati, usul va ob-havoga tayangan umumiy reja hamda maydon bo‘yicha suv hisobi.", accent="4E8A76")
add_card(slide, 0.72, 4.02, 5.75, 1.75, "Ob-havo", "Karmana (Navoiy) uchun joriy holat va 3 kunlik prognoz; internet bo‘lmasa saqlangan ma’lumot.", accent="5286A8")
add_card(slide, 6.78, 4.02, 5.75, 1.75, "Subsidiya va tarix", "Dasturlarni qidirish, mas’ul idora va hujjatlar haqida yo‘riqnoma; tahlillar qurilmada saqlanadi.", accent="C9851B")
add_text(slide, "Subsidiya bo‘limi ariza yubormaydi; shartlarni rasmiy manbadan qayta tekshirish kerak.",
         0.75, 6.16, 11.7, 0.34, 10, MUTED)

# Innovation and differentiation
slide = prs.slides.add_slide(blank)
add_box(slide, 0, 0, 13.333, 7.5, WHITE)
add_heading(slide, "04 / Innovatsiya va farqi", "Mahalliy ish jarayoniga moslangan yondashuv", 5)
add_box(slide, 0.75, 1.9, 5.65, 3.95, CREAM, rounded=True)
add_text(slide, "INNOVATSION YONDASHUV", 1.05, 2.22, 4.8, 0.3, 10, PRIMARY, True)
add_text(slide, "Offline AI +\nmaydonga mos hisob", 1.05, 2.75, 4.7, 1.0, 23, INK, True, "Aptos Display")
add_text(slide, "Kasallik rasmini telefonda tahlil qilishni, ekin va gektardan kelib chiqqan suv tejamkorligi hisobini, ob-havo tavsiyasini bitta sodda oqimga birlashtiradi.",
         1.05, 4.05, 4.75, 1.25, 13, MUTED)
add_box(slide, 6.75, 1.9, 5.8, 3.95, GREEN, rounded=True)
add_text(slide, "FARQLANISHI", 7.08, 2.22, 4.8, 0.3, 10, YELLOW, True)
diffs = [
    "O‘zbek tilidagi amaliy interfeys",
    "Kasallik tahlili qurilmaning o‘zida",
    "Maydonni kiritib, taxminiy suv tejashni ko‘rish",
    "Ob-havo, sug‘orish va subsidiya yo‘riqnomasi bir joyda",
]
for index, item in enumerate(diffs):
    y = 2.78 + index * 0.68
    add_box(slide, 7.1, y + 0.08, 0.12, 0.12, YELLOW, rounded=True)
    add_text(slide, item, 7.42, y, 4.65, 0.45, 13, WHITE, index == 0)
add_text(slide, "Farq alohida texnologiya borligida emas, funksiyalarni mahalliy va offline foydalanishga mos birlashtirishda.",
         0.78, 6.18, 11.8, 0.42, 12, MUTED, True)

# Water calculator
slide = prs.slides.add_slide(blank)
add_box(slide, 0, 0, 13.333, 7.5, GREEN)
add_heading(slide, "05 / Suv tejamkorligi", "Maydon kattalashsa, taxminiy tejam ham o‘zgaradi", 6, True)
add_text(slide, "Hisob tanlangan ekinning mavsumiy suv diapazoni va tomchilatib sug‘orishda ~45% kam sarf bo‘lishi haqidagi dasturiy taxminga asoslanadi.",
         0.76, 1.85, 7.2, 1.05, 15, "D0DFD2")
add_text(slide, "1 gektar g‘o‘za", 0.78, 3.38, 5.7, 0.45, 18, WHITE, True)
add_text(slide, "2 700–4 050 m³", 0.75, 4.02, 7.1, 0.9, 31, YELLOW, True, "Aptos Display")
add_text(slide, "bir mavsumda, yer ustidan sug‘orishga nisbatan taxminiy tejam", 0.8, 5.02, 6.5, 0.65, 13, "D0DFD2")
add_box(slide, 8.35, 2.05, 4.15, 3.5, "174B37", rounded=True)
add_text(slide, "HISOB MANTIG‘I", 8.72, 2.42, 3.3, 0.3, 10, YELLOW, True)
add_text(slide, "ekin me’yori\n× maydon\n× 45%", 8.72, 2.95, 3.2, 1.5, 21, WHITE, True, "Aptos Display")
add_text(slide, "Natija dala o‘lchovi emas. Tuproq, ob-havo, ekin bosqichi va tizim sozlamasi bo‘yicha pilotda tekshirilishi kerak.",
         0.78, 6.05, 11.45, 0.55, 11, "B8CCBC")

# Outcomes and validation
slide = prs.slides.add_slide(blank)
add_box(slide, 0, 0, 13.333, 7.5, CREAM)
add_heading(slide, "06 / Kutilayotgan natija", "Foydani sinovda o‘lchash va isbotlash", 7)
add_card(slide, 0.72, 1.85, 3.82, 2.0, "Tezroq yo‘naltirish", "Kasallik alomatlari bo‘yicha dastlabki taxmin va keyingi qadamni olish.", "01", PRIMARY)
add_card(slide, 4.76, 1.85, 3.82, 2.0, "Suv rejasini tushunish", "Maydon bo‘yicha suv miqdori va taxminiy tejamni solishtirish.", "02", "4E8A76")
add_card(slide, 8.8, 1.85, 3.82, 2.0, "Ma’lumotni jamlash", "Ob-havo, yo‘riqnoma va oldingi tahlillarni bir joyda ko‘rish.", "03", "C9851B")
add_text(slide, "Keyingi bosqich: fermerlar bilan dala sinovi", 0.75, 4.38, 8.1, 0.45, 19, INK, True, "Aptos Display")
metrics = [
    "Model taxminining agronom bahosi",
    "Tahlilga ketgan vaqt va qurilmadagi tezlik",
    "Hisoblangan va amalda sarflangan suv",
    "Foydalanuvchi tushunishi va qayta foydalanishi",
]
for index, item in enumerate(metrics):
    x = 0.8 + (index % 2) * 6.0
    y = 5.1 + (index // 2) * 0.62
    add_box(slide, x, y + 0.07, 0.12, 0.12, PRIMARY, rounded=True)
    add_text(slide, item, x + 0.28, y, 5.4, 0.35, 11, MUTED)
add_text(slide, "Hozirgi bosqich: funksional prototip. Natijalar va suv tejalishi bo‘yicha dala o‘lchovlari hali to‘planishi kerak.",
         0.78, 6.52, 11.6, 0.32, 10, PRIMARY, True)

prs.core_properties.title = "Dehqon AI — loyiha taqdimoti"
prs.core_properties.subject = "O‘zbek fermerlari uchun mobil qishloq xo‘jaligi yordamchisi"
prs.core_properties.author = "Dehqon AI"
prs.save(OUTPUT)
print(OUTPUT)